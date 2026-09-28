"""Offline research extraction, grounded questions, and room protocol checks."""

from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import game_server
from defense_session import DefenseSession
from project_files import ProjectFile
from question_generator import (
    AnsweredQuestion, CRITICAL_JUDGE, ETHICS_REVIEWER, IMPACT_REVIEWER,
    METHODOLOGY_REVIEWER, NextMoveDraft, PanelMove, QuestionDraft, QuestionGenerationError,
    RESEARCH_PANELISTS, generate_first_question, generate_next_move,
)
from research_files import combine_sources, read_research_files


def upload(name, data):
    return SimpleNamespace(name=name, getvalue=lambda: data)


def pdf_bytes(*texts, encrypted=False):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({
            NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 50 250 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    if encrypted:
        writer.encrypt('password')
    result = BytesIO()
    writer.write(result)
    return result.getvalue()


def docx_bytes():
    document = Document()
    document.add_paragraph('Planned interviews with students')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Consent'
    table.cell(0, 1).text = 'Anonymous responses'
    document.add_paragraph('Compare themes after interviews')
    result = BytesIO()
    document.save(result)
    return result.getvalue()


class FakeClient:
    def __init__(self, result):
        self.responses = self
        self.result = result
        self.request = None

    def parse(self, **kwargs):
        self.request = kwargs
        return SimpleNamespace(output_parsed=self.result)


class ResearchFilesTests(unittest.TestCase):
    def test_pdf_pages_and_docx_body_order_have_citation_locations(self):
        files, errors = read_research_files([
            upload('paper.pdf', pdf_bytes('Method pilot', 'Ethics consent')),
            upload('study.docx', docx_bytes()),
        ])
        self.assertFalse(errors)
        self.assertEqual(files[0].location_for(1), 'Page 1 · extracted line 1')
        self.assertEqual(files[0].location_for(2), 'Page 2 · extracted line 1')
        self.assertEqual(files[1].content.splitlines(), [
            'Planned interviews with students', 'Consent', 'Anonymous responses',
            'Compare themes after interviews',
        ])
        self.assertEqual(files[1].location_for(2), 'Table 1, row 1, cell 1')
        self.assertEqual(files[1].location_for(4), 'Paragraph 2')

    def test_unreadable_pages_encryption_and_limits_fail_before_defense(self):
        for data, expected in [
            (pdf_bytes('Readable', ''), 'page 2 has no readable text'),
            (pdf_bytes('Secret', encrypted=True), 'encrypted'),
            (b'not a pdf', 'could not be read'),
            (pdf_bytes(*(['Text'] * 101)), '100-page'),
        ]:
            with self.subTest(expected=expected):
                self.assertIn(expected, read_research_files([upload('paper.pdf', data)])[1][0])
        self.assertIn('10 MB', read_research_files([upload('paper.pdf', b'x' * 10_000_001)])[1][0])
        self.assertIn('600 KB', read_research_files([upload('paper.txt', b'a' * 600_001)])[1][0])
        paper = ProjectFile('paper.md', 'b' * 300_000, 'research_text')
        code = ProjectFile('code.py', 'a' * 300_001)
        self.assertIn('600 KB', combine_sources([code], [paper])[1][0])

    def test_research_questions_cite_extracted_paper_and_stage(self):
        paper = read_research_files([upload('paper.pdf', pdf_bytes('Planned student interviews'))])[0][0]
        code = ProjectFile('app.py', 'print("hello")')
        first_client = FakeClient(QuestionDraft(lead_in='', question='How will you recruit?', source_file=2, evidence_line=1))
        question = generate_first_question([code, paper], 'test-key', client=first_client,
                                           defense_type='mixed', research_stage='proposal')
        self.assertEqual(question.evidence_location, 'Page 1 · extracted line 1')
        self.assertEqual(question.evidence_text, 'Planned student interviews')
        self.assertIn('planned methods', first_client.request['input'][0]['content'])
        self.assertIn('research plan has a meaningful edge case', first_client.request['input'][0]['content'])
        self.assertIn('never assert that the event occurred', first_client.request['input'][0]['content'])
        self.assertIn('app.py', first_client.request['input'][1]['content'])
        self.assertIn('paper.pdf', first_client.request['input'][1]['content'])
        with self.assertRaisesRegex(QuestionGenerationError, 'invalid project file'):
            generate_first_question([code, paper], 'test-key', defense_type='mixed',
                client=FakeClient(QuestionDraft(lead_in='', question='How?', source_file=1, evidence_line=1)))
        session = DefenseSession.start(question, 'mixed', 'proposal')
        self.assertEqual(session.panelist_order, RESEARCH_PANELISTS)
        session.submit_answer('We will invite students who visit the office.')
        self.assertEqual(session.allowed_next_panelists, (METHODOLOGY_REVIEWER, ETHICS_REVIEWER))
        move = NextMoveDraft(action='ask', panelist=ETHICS_REVIEWER, lead_in='You will invite visitors.',
                             question='How will consent work?', source_file=2, evidence_line=1)
        client = FakeClient(move)
        result = generate_next_move([code, paper], 'test-key', history=session.answered_history(),
            allowed_panelists=session.allowed_next_panelists, may_complete=False,
            defense_type='mixed', research_stage='completed', client=client)
        self.assertEqual(result.panelist, ETHICS_REVIEWER)
        self.assertIn('reported results', client.request['input'][0]['content'])
        self.assertEqual(result.question.evidence_location, question.evidence_location)
        with self.assertRaisesRegex(QuestionGenerationError, 'invalid project file'):
            generate_next_move([code, paper], 'test-key', history=session.answered_history(),
                allowed_panelists=session.allowed_next_panelists, may_complete=False,
                defense_type='mixed', client=FakeClient(NextMoveDraft(
                    action='ask', panelist=ETHICS_REVIEWER, lead_in='', question='How?',
                    source_file=1, evidence_line=1)))
        session.apply_move(result)
        session.time_out_current()
        self.assertTrue(session.answered_history()[-1].timed_out)
        self.assertEqual(session.allowed_next_panelists, (ETHICS_REVIEWER, IMPACT_REVIEWER))

    def test_four_and_eight_research_turns(self):
        for followups in (False, True):
            with self.subTest(followups=followups):
                first = generate_first_question(
                    [ProjectFile('paper.md', 'Study proposal', 'research_text')], 'test-key',
                    defense_type='research', client=FakeClient(QuestionDraft(
                        lead_in='', question='How?', source_file=1, evidence_line=1)))
                session = DefenseSession.start(first, 'research')
                for index, role in enumerate(RESEARCH_PANELISTS):
                    session.submit_answer(f'Answer {index}')
                    if followups:
                        followup = session.allowed_next_panelists[0]
                        session.apply_move(PanelMove(followup, first))
                        session.time_out_current()
                    if role != CRITICAL_JUDGE:
                        session.apply_move(PanelMove(RESEARCH_PANELISTS[index + 1], first))
                    elif not followups:
                        session.apply_move(PanelMove(None, None))
                self.assertTrue(session.completed)
                self.assertEqual(len(session.turns), 8 if followups else 4)


class ResearchRoomTests(unittest.TestCase):
    def setUp(self):
        game_server.rooms.clear()
        self.client = TestClient(game_server.app)
        self.client.__enter__()
        self.addCleanup(lambda: self.client.__exit__(None, None, None))

    def test_research_room_creation_snapshot_and_mixed_requirements(self):
        response = self.client.post('/api/rooms', data={
            'host_name': 'Alex', 'defense_type': 'research', 'research_stage': 'proposal',
        }, files=[('research_files', ('paper.pdf', pdf_bytes('Planned pilot'), 'application/pdf'))])
        self.assertEqual(response.status_code, 201, response.text)
        host = response.json()
        room = game_server.rooms[host['room_code']]
        snapshot = room.snapshot(room.players[host['player_token']])
        self.assertEqual(snapshot['players'][0]['name'], 'Alex')
        self.assertEqual(snapshot['defense_type'], 'research')
        self.assertEqual(snapshot['research_stage'], 'proposal')
        self.assertEqual(snapshot['accepted_files'][0]['detail'], '1 page')
        guest = self.client.post(f"/api/rooms/{host['room_code']}/join", json={'name': 'Sam'}).json()
        with self.client.websocket_connect(f"/ws/{host['room_code']}") as socket:
            socket.send_json({'type': 'hello', 'token': guest['player_token']})
            joined = socket.receive_json()['state']
            self.assertEqual(joined['defense_type'], 'research')
            self.assertEqual(joined['accepted_files'], snapshot['accepted_files'])
            self.assertEqual(joined['players'][0]['name'], 'Alex')
        invalid = self.client.post('/api/rooms', data={'host_name': 'Alex', 'defense_type': 'mixed'},
            files=[('research_files', ('paper.md', b'Proposal', 'text/markdown'))])
        self.assertEqual(invalid.status_code, 422)

class ResearchTwoClientTests(unittest.TestCase):
    def setUp(self):
        game_server.rooms.clear()
        self.now = 1_000_000
        clock = patch.object(game_server, '_now_ms', side_effect=lambda: self.now)
        clock.start()
        self.addCleanup(clock.stop)
        from question_generator import SubmissionDecision
        interpretation_patch = patch.object(game_server, 'interpret_submission', return_value=SubmissionDecision('answer'))
        interpretation_patch.start()
        self.addCleanup(interpretation_patch.stop)
        self.client = TestClient(game_server.app)
        self.client.__enter__()
        self.addCleanup(lambda: self.client.__exit__(None, None, None))

    @staticmethod
    def until(socket, phase, turn_count=None):
        for _ in range(35):
            event = socket.receive_json()
            if event['type'] == 'snapshot':
                state = event['state']
                if state['phase'] == phase and (turn_count is None or len(state['turns']) == turn_count):
                    return state
            elif event['type'] == 'error':
                raise AssertionError(event['message'])
        raise AssertionError(f'Missing {phase} snapshot')

    def test_two_clients_complete_four_and_eight_research_turns(self):
        from question_generator import CoachingReport, GroundedQuestion
        for followups in (False, True):
            with self.subTest(followups=followups):
                game_server.rooms.clear()
                response = self.client.post('/api/rooms', data={
                    'host_name': 'Alex', 'defense_type': 'research', 'research_stage': 'proposal',
                }, files=[('research_files', ('paper.pdf', pdf_bytes('Planned student interviews'), 'application/pdf'))])
                self.assertEqual(response.status_code, 201, response.text)
                host = response.json()
                guest = self.client.post(f"/api/rooms/{host['room_code']}/join", json={'name': 'Sam'}).json()
                roles = [role for role in RESEARCH_PANELISTS for _ in range(2)] if followups else list(RESEARCH_PANELISTS)
                questions = [GroundedQuestion(f'How will turn {i + 1} work?', 'paper.pdf', 1,
                    'Planned student interviews', '', 'Page 1 · extracted line 1', 'research_pdf') for i in range(len(roles))]
                moves = [PanelMove(role, question) for role, question in zip(roles[1:], questions[1:])]
                if not followups:
                    moves.append(PanelMove(None, None))
                report = CoachingReport('Review complete.', [{'turn': 0, 'text': 'Specific plan.'}],
                                         [{'turn': 1, 'text': 'Clarify ethics.'}], 'Run a pilot.')
                with patch.object(game_server, 'generate_first_question', return_value=questions[0]), \
                     patch.object(game_server, 'generate_next_move', side_effect=moves), \
                     patch.object(game_server, 'generate_coaching_report', return_value=report):
                    with self.client.websocket_connect(f"/ws/{host['room_code']}") as host_ws, \
                         self.client.websocket_connect(f"/ws/{host['room_code']}") as guest_ws:
                        host_ws.send_json({'type': 'hello', 'token': host['player_token']})
                        self.until(host_ws, 'lobby')
                        guest_ws.send_json({'type': 'hello', 'token': guest['player_token']})
                        self.until(guest_ws, 'lobby')
                        host_ws.send_json({'type': 'start'})
                        for turn in range(len(roles)):
                            for ws in (host_ws, guest_ws):
                                state = self.until(ws, 'voting', turn + 1)
                                self.assertEqual(state['turns'][-1]['panelist'], roles[turn])
                                self.assertEqual(state['turns'][-1]['evidence_text'], 'Planned student interviews')
                                self.assertEqual(state['turns'][-1]['evidence_location'], 'Page 1 · extracted line 1')
                            host_ws.send_json({'type': 'cast_vote', 'seat': 0})
                            for ws in (host_ws, guest_ws):
                                voted = self.until(ws, 'voting', turn + 1)
                                self.assertEqual(voted['vote_counts'].get('0'), 1)
                            self.now += game_server.VOTE_MS + 1
                            room = game_server.rooms[host['room_code']]
                            self.client.portal.call(game_server._expire_deadline, room)
                            for ws in (host_ws, guest_ws):
                                state = self.until(ws, 'question', turn + 1)
                                self.assertEqual(state['selected_seat'], 0)
                            if followups and turn == 1:
                                self.now += game_server.ANSWER_MS + 1
                                self.client.portal.call(game_server._expire_deadline, room)
                            else:
                                host_ws.send_json({'type': 'submit_answer', 'turn': turn,
                                                   'answer': f'Answer {turn + 1}'})
                        for ws in (host_ws, guest_ws):
                            state = self.until(ws, 'complete', len(roles))
                            self.assertEqual(state['turns'][1]['timed_out'], followups)
                            state = self.until(ws, 'complete', len(roles)) if state['feedback_status'] != 'ready' else state
                            self.assertEqual(state['feedback']['summary'], 'Review complete.')
                        # Reconnected teammates receive the same defense type, transcript, and coaching.
                    with self.client.websocket_connect(f"/ws/{host['room_code']}") as reconnect:
                        reconnect.send_json({'type': 'hello', 'token': guest['player_token']})
                        state = self.until(reconnect, 'complete', len(roles))
                        self.assertEqual(state['defense_type'], 'research')
                        self.assertEqual(len(state['turns']), len(roles))

    def test_research_generation_retry_keeps_answer_and_reconnect(self):
        from question_generator import GroundedQuestion
        response = self.client.post('/api/rooms', data={
            'host_name': 'Alex', 'defense_type': 'mixed', 'research_stage': 'infer',
        }, files=[
            ('files', ('app.py', b'USE_PILOT = True\n', 'text/x-python')),
            ('research_files', ('study.md', b'Pilot interviews are planned.\n', 'text/markdown')),
        ])
        self.assertEqual(response.status_code, 201, response.text)
        host = response.json()
        guest = self.client.post(f"/api/rooms/{host['room_code']}/join", json={'name': 'Sam'}).json()
        first = GroundedQuestion('How will the pilot work?', 'study.md', 1,
                                 'Pilot interviews are planned.', evidence_kind='research_text', evidence_location='Line 1')
        second = GroundedQuestion('How will consent work?', 'study.md', 1,
                                  'Pilot interviews are planned.', evidence_kind='research_text', evidence_location='Line 1')
        with patch.object(game_server, 'generate_first_question', return_value=first), \
             patch.object(game_server, 'generate_next_move', side_effect=[
                QuestionGenerationError('The AI cited an invalid source line. Please try again.'),
                PanelMove(ETHICS_REVIEWER, second)]):
            with self.client.websocket_connect(f"/ws/{host['room_code']}") as host_ws:
                host_ws.send_json({'type': 'hello', 'token': host['player_token']})
                self.until(host_ws, 'lobby')
                host_ws.send_json({'type': 'start'})
                self.until(host_ws, 'voting', 1)
                host_ws.send_json({'type': 'cast_vote', 'seat': 0})
                self.now += game_server.VOTE_MS + 1
                self.client.portal.call(game_server._expire_deadline, game_server.rooms[host['room_code']])
                self.until(host_ws, 'question', 1)
                host_ws.send_json({'type': 'submit_answer', 'turn': 0, 'answer': 'We will recruit volunteers.'})
                state = self.until(host_ws, 'retry', 1)
                self.assertEqual(state['turns'][0]['answer'], 'We will recruit volunteers.')
                with self.client.websocket_connect(f"/ws/{host['room_code']}") as guest_ws:
                    guest_ws.send_json({'type': 'hello', 'token': guest['player_token']})
                    reconnected = self.until(guest_ws, 'retry', 1)
                    self.assertEqual(reconnected['turns'][0]['answer'], state['turns'][0]['answer'])
                    host_ws.send_json({'type': 'retry'})
                    for ws in (host_ws, guest_ws):
                        advanced = self.until(ws, 'voting', 2)
                        self.assertEqual(advanced['turns'][1]['panelist'], ETHICS_REVIEWER)
                        self.assertEqual(advanced['turns'][0]['answer'], 'We will recruit volunteers.')
