from offline_accounts import authenticate
"""Offline paper-map checks; all AI calls mocked, no paid requests."""

import asyncio
from io import BytesIO
import os
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from pydantic import ValidationError
from streamlit.testing.v1 import AppTest

import game_server
from project_files import ProjectFile
from question_generator import GroundedQuestion, QuestionGenerationError, RESEARCH_PANELISTS
from research_files import read_research_files
from research_plan import generate_research_plan, validate_question_budget
from test_research_defense import pdf_bytes, upload


PAPER = "# Purpose\nStudy queue waiting times.\n# Methods\nRecruit volunteers for a two-week pilot.\n# Ethics\nObtain consent and anonymize responses.\n# Impact\nCompare students' wait times.\n# Limitations\nA single campus may not represent other campuses.\n# References\nAn existing queue study.\n"


def draft():
    return {"topics": [
        {"title": title, "objective": "Discuss " + title.lower() + " using the stated pilot plan.",
         "panelist": role, "references": [{"source_file": 1, "evidence_line": line}],
         "gaps": ["The participant count is not stated."] if index == 0 else []}
        for index, (title, role, line) in enumerate(zip(
            ("Study design", "Participant safeguards", "Expected impact", "Limitations"),
            RESEARCH_PANELISTS, (4, 6, 8, 10)))],
        "uncertainties": ["No completed results are stated in the proposal."]}


class FakeClient:
    def __init__(self, output):
        self.output, self.calls = output, []
        self.responses = self

    def parse(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(output_parsed=self.output)


def fixture_plan():
    return generate_research_plan([ProjectFile("study.md", PAPER, "research_text")],
                                  "offline-key", research_stage="proposal", client=FakeClient(draft()))


class ResearchPlanGenerationTests(unittest.TestCase):
    def setUp(self):
        self.files = [ProjectFile("study.md", PAPER, "research_text")]

    def test_single_request_maps_order_roles_gaps_and_owned_excerpts(self):
        output = draft()
        output["topics"][0]["references"][0]["evidence_text"] = "UNTRUSTED AI EXCERPT"
        client = FakeClient(output)
        plan = generate_research_plan(self.files, "offline-key", research_stage="proposal", client=client)
        self.assertEqual(len(client.calls), 1)
        request = client.calls[0]
        self.assertEqual(request["model"], "gpt-6-luna")
        self.assertFalse(request["store"])
        self.assertEqual(request["reasoning"], {"effort": "low"})
        prompt = request["input"][0]["content"]
        for phrase in ("Cover every substantive section", "omit references-only", "not questions or answers",
                       "untrusted project data", "proposal", "do not imply results already exist"):
            self.assertIn(phrase, prompt)
        self.assertIn("An existing queue study.", request["input"][1]["content"])
        self.assertEqual([topic.panelist for topic in plan.topics], list(RESEARCH_PANELISTS))
        self.assertEqual([topic.id for topic in plan.topics], [f"topic-{i}" for i in range(1, 5)])
        self.assertEqual(plan.topics[0].references[0].evidence_text, PAPER.splitlines()[3])
        self.assertEqual(plan.topics[0].references[0].evidence_location, "Line 4")
        self.assertNotIn("UNTRUSTED AI EXCERPT", str(plan.snapshot()))
        self.assertEqual(plan.suggested_budget, 8)
        self.assertTrue(plan.topics[0].gaps)
        self.assertTrue(plan.uncertainties)

    def test_all_stages_and_model_override_reach_prompt(self):
        for stage, phrase in (("proposal", "planned methods and feasibility"), ("completed", "reported results"),
                              ("infer", "if unclear, ask for clarification")):
            with self.subTest(stage=stage):
                client = FakeClient(draft())
                generate_research_plan(self.files, "offline-key", "configured-model", research_stage=stage, client=client)
                self.assertEqual(client.calls[0]["model"], "configured-model")
                self.assertIn(phrase, client.calls[0]["input"][0]["content"])

    def test_mixed_map_includes_code_and_paper_references(self):
        output = draft()
        output["topics"][0]["references"].append({"source_file": 2, "evidence_line": 1})
        client = FakeClient(output)
        plan = generate_research_plan(self.files + [ProjectFile("queue.py", "DATABASE = 'queue.db'\n")],
                                     "offline-key", defense_type="mixed", client=client)
        self.assertEqual([r.filename for r in plan.topics[0].references], ["study.md", "queue.py"])
        self.assertEqual(plan.topics[0].references[1].evidence_text, "DATABASE = 'queue.db'")
        self.assertIn("implementation choices or discrepancies", client.calls[0]["input"][0]["content"])
        self.assertIn("DATABASE = 'queue.db'", client.calls[0]["input"][1]["content"])

    def test_pdf_and_docx_references_keep_extracted_locations(self):
        document = Document()
        document.add_paragraph("Pilot methodology")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Consent before interviews"
        stream = BytesIO()
        document.save(stream)
        files, errors = read_research_files([
            upload("paper.pdf", pdf_bytes("First page purpose", "Second page methods")),
            upload("appendix.docx", stream.getvalue()),
        ])
        self.assertFalse(errors)
        refs = [{"source_file": 1, "evidence_line": 2}, {"source_file": 2, "evidence_line": 2}]
        output = draft()
        output["topics"] = [{**output["topics"][0], "references": refs}]
        plan = generate_research_plan(files, "offline-key", client=FakeClient(output))
        self.assertEqual(plan.topics[0].references[0].evidence_location, "Page 2 · extracted line 1")
        self.assertIn("Table 1", plan.topics[0].references[1].evidence_location)
        self.assertEqual(plan.topics[0].references[1].evidence_text, "Consent before interviews")
        self.assertEqual(plan.topics[0].references[0].evidence_before, "")

    def test_invalid_empty_blank_or_unowned_references_are_rejected(self):
        for refs in ([], [{"source_file": 99, "evidence_line": 1}], [{"source_file": 1, "evidence_line": 99}],
                     [{"source_file": 1, "evidence_line": 1}] * 2):
            with self.subTest(refs=refs):
                output = draft()
                output["topics"][0]["references"] = refs
                with self.assertRaises(QuestionGenerationError):
                    generate_research_plan(self.files, "offline-key", client=FakeClient(output))
        output = draft()
        output["topics"][0]["references"] = [{"source_file": 1, "evidence_line": 2}]
        with self.assertRaisesRegex(QuestionGenerationError, "invalid source line"):
            generate_research_plan([ProjectFile("study.md", "Heading\n\n", "research_text")], "offline-key", client=FakeClient(output))

    def test_code_only_topic_is_invalid_even_in_mixed_mode(self):
        output = draft()
        output["topics"][0]["references"] = [{"source_file": 2, "evidence_line": 1}]
        with self.assertRaisesRegex(QuestionGenerationError, "must cite an uploaded paper"):
            generate_research_plan(self.files + [ProjectFile("app.py", "print('demo')")], "offline-key",
                                   defense_type="mixed", client=FakeClient(output))

    def test_bad_roles_noninteger_lines_duplicates_and_text_limits_fail(self):
        for changes in ({"panelist": "Security Reviewer"}, {"title": ""}, {"objective": "x" * 601},
                        {"references": [{"source_file": True, "evidence_line": 1}]}):
            with self.subTest(changes=changes):
                output = draft()
                output["topics"][0].update(changes)
                with self.assertRaises((QuestionGenerationError, ValidationError)):
                    generate_research_plan(self.files, "offline-key", client=FakeClient(output))
        output = draft()
        output["topics"][1]["title"] = output["topics"][0]["title"]
        with self.assertRaisesRegex(QuestionGenerationError, "duplicate topics"):
            generate_research_plan(self.files, "offline-key", client=FakeClient(output))

    def test_missing_credentials_paper_and_combined_limits_fail_without_call(self):
        for files, key in ((self.files, None), (self.files, "bad key"), ([ProjectFile("app.py", "x")], "offline-key"),
                           ([ProjectFile("study.md", "a" * 600001, "research_text")], "offline-key")):
            client = FakeClient(draft())
            with self.assertRaises(QuestionGenerationError):
                generate_research_plan(files, key, client=client)
            self.assertFalse(client.calls)
        with self.assertRaises(ValueError):
            generate_research_plan(self.files, "offline-key", defense_type="code", client=FakeClient(draft()))

    def test_no_parsed_map_and_budget_bounds(self):
        with self.assertRaisesRegex(QuestionGenerationError, "no research map"):
            generate_research_plan(self.files, "offline-key", client=FakeClient(None))
        for count, expected in ((1, 4), (9, 18), (51, 100)):
            output = draft()
            output["topics"] = [{**output["topics"][0], "title": f"Topic {i}"} for i in range(count)]
            plan = generate_research_plan(self.files, "offline-key", client=FakeClient(output))
            self.assertEqual(plan.suggested_budget, expected)
        for value in (True, 4.0, "8", 3, 101, None):
            with self.assertRaises(ValueError):
                validate_question_budget(value)
        self.assertEqual(validate_question_budget(100), 100)


class ResearchPlanRoomTests(unittest.TestCase):
    def setUp(self):
        game_server.rooms.clear()
        env = patch.dict(os.environ, {"GAME_HOST_PASSCODE": "offline-passcode", "OPENAI_API_KEY": "offline-key"})
        env.start()
        self.addCleanup(env.stop)
        authenticate(self)
        self.client = TestClient(game_server.app)
        self.client.__enter__()
        self.addCleanup(lambda: self.client.__exit__(None, None, None))
        authenticate(self, self.client)
        self.plan = fixture_plan()

    def create(self, mode="research"):
        files = [("research_files", ("study.md", PAPER.encode(), "text/markdown"))]
        if mode == "mixed":
            files.append(("files", ("queue.py", b"DATABASE='queue.db'", "text/x-python")))
        result = self.client.post("/api/rooms", data={"host_name": "Host", "host_passcode": "offline-passcode",
            "defense_type": mode, "research_stage": "proposal"}, files=files)
        self.assertEqual(result.status_code, 201, result.text)
        return result.json()

    def hello(self, socket, token):
        socket.send_json({"type": "hello", "token": token})
        return socket.receive_json()["state"]

    def until(self, socket, **conditions):
        for _ in range(30):
            event = socket.receive_json()
            if event["type"] == "error":
                self.fail(event["message"])
            state = event["state"]
            if all(state.get(key) == value for key, value in conditions.items()):
                return state
        self.fail("Missing map state")

    def test_two_clients_prepare_confirm_and_reconnect_without_extra_call(self):
        for mode in ("research", "mixed"):
            with self.subTest(mode=mode), patch("game_server.generate_research_plan", return_value=self.plan) as generator, \
                 patch("game_server.generate_first_question") as question_call:
                host = self.create(mode)
                guest = self.client.post(f"/api/rooms/{host['room_code']}/join", json={"name": "Guest"}).json()
                with self.client.websocket_connect(f"/ws/{host['room_code']}") as h, \
                     self.client.websocket_connect(f"/ws/{host['room_code']}") as g:
                    self.hello(h, host["player_token"])
                    self.hello(g, guest["player_token"])
                    h.send_json({"type": "prepare_research_plan"})
                    for socket in (h, g):
                        state = self.until(socket, research_planning_status="ready")
                        self.assertEqual(state["phase"], "lobby")
                        self.assertEqual(state["turns"], [])
                        self.assertEqual(state["research_plan"]["topics"][0]["references"][0]["evidence_text"], PAPER.splitlines()[3])
                    for action in ("prepare_research_plan", "retry_research_plan", "approve_research_plan"):
                        g.send_json({"type": action})
                        self.assertIn("Only the host", g.receive_json()["message"])
                    h.send_json({"type": "approve_research_plan", "plan_id": self.plan.id, "question_budget": 20})
                    for socket in (h, g):
                        state = self.until(socket, research_plan_approved=True)
                        self.assertEqual(state["research_budget_preview"], 20)
                    self.assertEqual(generator.call_count, 1)
                    self.assertEqual(generator.call_args.kwargs["defense_type"], mode)
                    self.assertEqual(generator.call_args.kwargs["research_stage"], "proposal")
                    if mode == "mixed":
                        self.assertEqual(len(generator.call_args.args[0]), 2)
                    question_call.assert_not_called()
                with self.client.websocket_connect(f"/ws/{host['room_code']}") as restored:
                    state = self.hello(restored, guest["player_token"])
                    self.assertTrue(state["research_plan_approved"])
                    self.assertEqual(state["research_budget_preview"], 20)
                    self.assertEqual(state["research_plan"]["id"], self.plan.id)
                self.assertEqual(generator.call_count, 1)

    def test_failure_retry_and_stale_or_invalid_approval_preserve_uploads(self):
        host = self.create()
        with patch("game_server.generate_research_plan", side_effect=[QuestionGenerationError("The AI cited an invalid source line."), self.plan]) as generator, \
             patch("game_server.generate_first_question") as question_call, \
             self.client.websocket_connect(f"/ws/{host['room_code']}") as h:
            self.hello(h, host["player_token"])
            h.send_json({"type": "prepare_research_plan"})
            state = self.until(h, research_planning_status="failed")
            self.assertEqual(state["files"], ["study.md"])
            self.assertIsNone(state["research_plan"])
            self.assertEqual(state["phase"], "lobby")
            h.send_json({"type": "retry_research_plan"})
            self.until(h, research_planning_status="ready")
            for plan_id, value in (("stale-plan", 8), (self.plan.id, True), (self.plan.id, 101)):
                h.send_json({"type": "approve_research_plan", "plan_id": plan_id, "question_budget": value})
                self.assertEqual(h.receive_json()["type"], "error")
            self.assertFalse(game_server.rooms[host["room_code"]].research_plan_approved)
            self.assertEqual(generator.call_count, 2)
            question_call.assert_not_called()

    def test_confirmed_budget_starts_the_research_policy(self):
        host = self.create()
        room = game_server.rooms[host["room_code"]]
        room.research_plan = self.plan
        room.research_planning_status = "ready"
        room.research_plan_approved = True
        room.research_budget_preview = 20
        q = GroundedQuestion("Explain the pilot design.", "study.md", 4, PAPER.splitlines()[3])
        with patch("game_server.generate_research_move", return_value=__import__("question_generator").ResearchMove(RESEARCH_PANELISTS[0], q, "topic-1", False, None)), \
             self.client.websocket_connect(f"/ws/{host['room_code']}") as h:
            self.hello(h, host["player_token"])
            h.send_json({"type": "start", "plan_id": self.plan.id, "question_budget": 20})
            state = self.until(h, phase="voting")
            self.assertEqual(state["research_budget_preview"], 20)
            self.assertEqual(len(state["turns"]), 1)
            room.defense.submit_answer("The pilot compares wait times.")
            self.assertEqual(room.defense.question_budget, 20)
            self.assertEqual(set(room.defense.allowed_next_panelists), set(RESEARCH_PANELISTS[:2]))
            h.send_json({"type": "prepare_research_plan"})
            self.assertIn("before starting", h.receive_json()["message"])


class ResearchPlanRaceTests(unittest.IsolatedAsyncioTestCase):
    async def test_start_during_planning_rejected_and_restart_discards_late_map(self):
        authenticate(self, unit_server=game_server)
        host = game_server.Player("host", "Host", 0, True)
        room = game_server.Room("OFFLINE", [ProjectFile("study.md", PAPER, "research_text")], {host.token: host}, "research")
        room.research_planning_status = "planning"
        room.research_plan_generation_id = 1
        events = []
        socket = SimpleNamespace(send_json=lambda event: capture(event))
        async def capture(event):
            events.append(event)
        await game_server._handle_action(room, host, socket, {"type": "start"})
        self.assertIn("Wait for", events[-1]["message"])
        with patch("game_server._schedule_research_plan") as schedule:
            await game_server._handle_action(room, host, socket, {"type": "prepare_research_plan"})
            self.assertIn("already being prepared", events[-1]["message"])
            schedule.assert_not_called()
        plan = fixture_plan()
        started, release = asyncio.Event(), asyncio.Event()
        async def delayed(*args, **kwargs):
            started.set()
            await release.wait()
            return plan
        with patch("game_server.asyncio.to_thread", side_effect=delayed), patch("game_server._schedule_generation"):
            task = asyncio.create_task(game_server._prepare_research_plan(room, 1))
            await started.wait()
            await game_server._handle_action(room, host, socket, {"type": "restart"})
            release.set()
            await task
        self.assertEqual(room.phase, "lobby")
        self.assertEqual(room.research_planning_status, "none")
        self.assertIsNone(room.research_plan)


class ResearchPlanStreamlitTests(unittest.TestCase):
    def test_prepare_retry_approve_and_material_change_reset_without_question(self):
        plan = fixture_plan()
        with patch.dict(os.environ, {"OPENAI_API_KEY": "offline-key"}), \
             patch("research_plan.generate_research_plan", side_effect=[QuestionGenerationError("Map failed."), plan]) as generator, \
             patch("question_generator.generate_first_question") as question_call:
            app = AppTest.from_file("app.py").run()
            app.selectbox[0].select("Research paper").run()
            app.file_uploader[0].set_value([("study.md", PAPER.encode(), "text/markdown")]).run()
            next(b for b in app.button if b.label == "Prepare defense").click().run()
            self.assertFalse(app.exception)
            self.assertTrue(any("Map failed." in error.value for error in app.error))
            next(b for b in app.button if b.label == "Retry research map").click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.session_state["research_plan"].id, plan.id)
            app.number_input[0].set_value(20).run()
            next(b for b in app.button if b.label == "Confirm question budget").click().run()
            self.assertTrue(app.session_state["research_plan_approved"])
            self.assertEqual(app.session_state["research_confirmed_budget"], 20)
            question_call.assert_not_called()
            self.assertEqual(generator.call_count, 2)
            app.file_uploader[0].set_value([("study.md", (PAPER + "New section\n").encode(), "text/markdown")]).run()
            self.assertNotIn("research_plan", app.session_state)
            self.assertNotIn("research_budget_preview", app.session_state)
            self.assertFalse(app.exception)


if __name__ == "__main__":
    unittest.main()
